import io
from datetime import date

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill

from gastos.periodos import nome_do_mes

FORMATO_MOEDA = '"R$" #,##0.00'
FORMATO_DATA = "DD/MM/YYYY"
CABECALHO_FONTE = Font(bold=True, color="FFFFFF")
CABECALHO_FUNDO = PatternFill("solid", fgColor="1A7F37")


def _cabecalho(aba, titulos, larguras):
    aba.append(titulos)
    for celula in aba[1]:
        celula.font = CABECALHO_FONTE
        celula.fill = CABECALHO_FUNDO
        celula.alignment = Alignment(horizontal="center")
    for indice, largura in enumerate(larguras, start=1):
        aba.column_dimensions[chr(64 + indice)].width = largura
    aba.freeze_panes = "A2"


def _aba_do_mes(planilha, mes, gastos):
    aba = planilha.create_sheet(title=mes)
    _cabecalho(aba, ["Data", "Descrição", "Categoria", "Valor", "Fixo"], [13, 34, 18, 16, 8])

    for gasto in sorted(gastos, key=lambda g: (g["data"], g["id"])):
        aba.append(
            [
                date.fromisoformat(gasto["data"]),
                gasto["descricao"],
                gasto["categoria"],
                gasto["valor"],
                "Sim" if gasto["fixo"] else "Não",
            ]
        )
        linha = aba.max_row
        aba.cell(linha, 1).number_format = FORMATO_DATA
        aba.cell(linha, 4).number_format = FORMATO_MOEDA

    if gastos:
        ultima = aba.max_row
        aba.auto_filter.ref = f"A1:E{ultima}"
        total = ultima + 1
        aba.cell(total, 3, "Total").font = Font(bold=True)
        aba.cell(total, 4, f"=SUM(D2:D{ultima})").font = Font(bold=True)
        aba.cell(total, 4).number_format = FORMATO_MOEDA


def _aba_resumo(planilha, resumo):
    aba = planilha.active
    aba.title = "Resumo"
    _cabecalho(aba, ["Mês", "Total gasto", "Gastos fixos", "Gastos variáveis", "Nº de gastos", "Meta"], [20, 16, 16, 18, 13, 16])

    for item in resumo:
        aba.append(
            [
                nome_do_mes(item["mes"]),
                item["total"],
                item["fixos"],
                item["total"] - item["fixos"],
                item["quantidade"],
                item["meta"],
            ]
        )
        linha = aba.max_row
        for coluna in (2, 3, 4, 6):
            aba.cell(linha, coluna).number_format = FORMATO_MOEDA

    if resumo:
        ultima = aba.max_row
        aba.cell(ultima + 1, 1, "Total").font = Font(bold=True)
        for coluna in (2, 3, 4):
            letra = chr(64 + coluna)
            celula = aba.cell(ultima + 1, coluna, f"=SUM({letra}2:{letra}{ultima})")
            celula.font = Font(bold=True)
            celula.number_format = FORMATO_MOEDA

        grafico = BarChart()
        grafico.title = "Total gasto por mês"
        grafico.legend = None
        grafico.add_data(Reference(aba, min_col=2, min_row=1, max_row=ultima), titles_from_data=True)
        grafico.set_categories(Reference(aba, min_col=1, min_row=2, max_row=ultima))
        grafico.height, grafico.width = 8, 16
        aba.add_chart(grafico, "H2")


def gerar_xlsx(gastos_por_mes, resumo=None):
    """Monta a planilha e devolve os bytes do arquivo .xlsx.

    gastos_por_mes: dicionário {'AAAA-MM': [gastos]} na ordem em que as abas devem aparecer.
    resumo: lista de dicionários (mes, total, fixos, quantidade, meta). Se vier, cria a aba "Resumo" primeiro.
    """
    planilha = Workbook()
    if resumo is not None:
        _aba_resumo(planilha, resumo)
    else:
        planilha.remove(planilha.active)

    for mes, gastos in gastos_por_mes.items():
        _aba_do_mes(planilha, mes, gastos)

    buffer = io.BytesIO()
    planilha.save(buffer)
    return buffer.getvalue()
