import io
from datetime import date

from openpyxl import load_workbook

from gastos.excel import gerar_xlsx


def gasto(id, data, descricao, categoria, valor, fixo=0):
    return {"id": id, "data": data, "descricao": descricao, "categoria": categoria, "valor": valor, "fixo": fixo}


def abrir(bytes_xlsx):
    return load_workbook(io.BytesIO(bytes_xlsx))


def test_uma_aba_por_mes_com_valores_reais():
    dados = {"2026-10": [gasto(2, "2026-10-12", "Cinema", "lazer", 48.9), gasto(1, "2026-10-05", "Aluguel", "moradia", 900.0, fixo=1)]}
    planilha = abrir(gerar_xlsx(dados))
    assert planilha.sheetnames == ["2026-10"]
    aba = planilha["2026-10"]
    assert [c.value for c in aba[1]] == ["Data", "Descrição", "Categoria", "Valor", "Fixo"]
    # ordenado por data crescente, com data e valor como tipos de verdade (não texto)
    assert aba["A2"].value.date() == date(2026, 10, 5)
    assert aba["B2"].value == "Aluguel"
    assert aba["D2"].value == 900.0
    assert aba["E2"].value == "Sim"
    assert aba["E3"].value == "Não"
    assert "R$" in aba["D2"].number_format


def test_linha_de_total_usa_formula_de_soma():
    dados = {"2026-10": [gasto(1, "2026-10-05", "A", "x", 10.0), gasto(2, "2026-10-06", "B", "x", 20.0)]}
    aba = abrir(gerar_xlsx(dados))["2026-10"]
    assert aba["C4"].value == "Total"
    assert aba["D4"].value == "=SUM(D2:D3)"


def test_aba_resumo_vem_primeiro_com_todos_os_meses():
    dados = {
        "2026-09": [gasto(1, "2026-09-05", "Aluguel", "moradia", 900.0, fixo=1), gasto(2, "2026-09-08", "Mercado", "alimentação", 400.0)],
        "2026-10": [gasto(3, "2026-10-05", "Aluguel", "moradia", 900.0, fixo=1)],
    }
    resumo = [
        {"mes": "2026-09", "total": 1300.0, "fixos": 900.0, "quantidade": 2, "meta": 1800.0},
        {"mes": "2026-10", "total": 900.0, "fixos": 900.0, "quantidade": 1, "meta": None},
    ]
    planilha = abrir(gerar_xlsx(dados, resumo))
    assert planilha.sheetnames == ["Resumo", "2026-09", "2026-10"]
    aba = planilha["Resumo"]
    assert aba["A2"].value == "setembro de 2026"
    assert (aba["B2"].value, aba["C2"].value, aba["D2"].value, aba["E2"].value, aba["F2"].value) == (1300.0, 900.0, 400.0, 2, 1800.0)
    assert aba["F3"].value is None
    assert aba["B4"].value == "=SUM(B2:B3)"


def test_sem_gastos_gera_arquivo_valido():
    assert abrir(gerar_xlsx({}, [])).sheetnames == ["Resumo"]
    assert abrir(gerar_xlsx({"2026-10": []})).sheetnames == ["2026-10"]
